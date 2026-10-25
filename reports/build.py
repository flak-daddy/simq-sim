import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

d=json.load(open('data.json')); T=d['trades']; W=d['weeks']
F='Arial'
BLUE=Font(name=F,color='0000FF',size=10); BLK=Font(name=F,size=10); BOLD=Font(name=F,size=10,bold=True)
GRN=Font(name=F,color='008000',size=10)
HDR=Font(name=F,size=10,bold=True,color='FFFFFF'); HFILL=PatternFill('solid',fgColor='1F3864')
SUB=PatternFill('solid',fgColor='D9E1F2'); YEL=PatternFill('solid',fgColor='FFFF00')
thin=Side(style='thin',color='BFBFBF'); BOX=Border(bottom=thin)
CUR='$#,##0.00;($#,##0.00);"-"'; CUR0='$#,##0;($#,##0);"-"'; PCT='0.0%;(0.0%);"-"'; PCT2='0.00%;(0.00%);"-"'
PH={'open':'Mon','mid':'Wed'}
OUT={'expire':'Expired worthless','btc':'Bought to close','assign':'Assigned','called':'Called away'}

wb=Workbook()
# ---------------- Assumptions on Summary sheet (created first so refs are stable)
S=wb.active; S.title='Summary'
# ---------------- Trades sheet
X=wb.create_sheet('Trades')
th=['Trade #','Week','Opened','Type','Strike','SimQ at sale','Cushion from spot','Premium / share','Premium in ($)','Collateral or cover','Premium % of collateral','Outcome','Buy-back / share','Buy-back out ($)','Net option P&L ($)','SimQ Friday close','Share result']
X.append(th)
for i,t in enumerate(T):
    r=i+2; wk=W[t['w']-1]
    put=t['kind']=='P'
    share=''
    if t['out']=='assign': share=f"Bought 100 at ${t['k']:g}"
    if t['out']=='called': share=f"Sold 100 at ${t['k']:g}"
    row=[t['id'],t['w'],PH[t['ph']],'Put' if put else 'Call',t['k'],t['spot'],
         f'=(F{r}-E{r})/F{r}' if put else f'=(E{r}-F{r})/F{r}',
         t['fill'],f'=H{r}*100',
         f'=E{r}*100' if put else '100 shares',
         f'=I{r}/J{r}' if put else f'=I{r}/(F{r}*100)',
         OUT[t['out']],
         t.get('close',0),f'=M{r}*100',f'=I{r}-N{r}',wk['fri'],share]
    X.append(row)
    for c in range(1,18):
        cell=X.cell(r,c); cell.font=BLK
    for c in (1,2,5,6,8,13,16): X.cell(r,c).font=BLUE
    for c,fmt in ((5,'0.0#'),(6,'0.00'),(7,PCT),(8,'0.00'),(9,CUR),(10,CUR0),(11,PCT2),(13,'0.00'),(14,CUR),(15,CUR),(16,'0.00')):
        X.cell(r,c).number_format=fmt
nT=len(T)+1
tr=nT+2
X.cell(tr,1,'Totals').font=BOLD
for c in (9,14,15):
    X.cell(tr,c,f'=SUM({L(c)}2:{L(c)}{nT})').font=BOLD; X.cell(tr,c).number_format=CUR
X.cell(tr+2,1,'Blue = recorded trade data (fills, strikes, prices). Black = formulas. Source: SimQ Run 2 trade log (run-002-state.json), simulated paper trading.').font=Font(name=F,size=9,italic=True)
X.cell(tr+3,1,'Fills: sold at the bid, bought back at the ask. Cushion = how far the strike sat from SimQ when sold (below spot for puts, above spot for calls).').font=Font(name=F,size=9,italic=True)
widths=[8,7,8,7,8,11,11,11,12,14,12,18,12,13,13,12,18]
for i,w in enumerate(widths,1): X.column_dimensions[L(i)].width=w
X.freeze_panes='C2'

# ---------------- Weekly sheet
Y=wb.create_sheet('Weekly')
wh=['Week','Regime (revealed Fri)','SimQ Mon open','SimQ Fri close','SimQ week %','Actions taken','Trades opened','Premium in ($)','Buy-backs out ($)','Options net ($)','Share P&L ($)','Interest ($)','Account start ($)','Account end ($)','Week P&L ($)','Week %','Hit 1%?','Cumulative return','1% weekly target path ($)','Ahead / behind target ($)','Shares held Fri','Buy-and-hold value ($)','Cash-only value ($)']
Y.append(wh)
held=0
for i,w in enumerate(W):
    r=i+2; n=w['w']
    acts=[]
    tw=[t for t in T if t['w']==n]
    for ph in ('open','mid'):
        if ph in w['passes']: acts.append(f"{PH[ph]}: pass/hold")
        for t in tw:
            if t['out']=='btc' and t['close_ph']==ph: acts.append(f"{PH[ph]}: bought back {t['kind']}{t['k']:g} @{t['close']:.2f}")
        for t in tw:
            if t['ph']==ph: acts.append(f"{PH[ph]}: sold {t['kind']}{t['k']:g} @{t['fill']:.2f}")
    for t in tw:
        if t['out']!='btc': acts.append(f"Fri: {t['kind']}{t['k']:g} {OUT[t['out']].lower()}")
    for t in tw:
        if t['out']=='assign': held+=100
        if t['out']=='called': held-=100
    row=[n,w['regime'],w['open'],w['fri'],f'=D{r}/C{r}-1','; '.join(acts),
         f'=COUNTIF(Trades!$B$2:$B${nT},A{r})',
         f'=SUMIF(Trades!$B$2:$B${nT},A{r},Trades!$I$2:$I${nT})',
         f'=SUMIF(Trades!$B$2:$B${nT},A{r},Trades!$N$2:$N${nT})',
         f'=H{r}-I{r}', w['shares'], w['interest'],
         '=Summary!$C$5' if n==1 else f'=N{r-1}',
         f'=M{r}+J{r}+K{r}+L{r}', f'=N{r}-M{r}', f'=O{r}/M{r}',
         f'=IF(P{r}>=Summary!$C$7,"Yes","No")',
         f'=N{r}/Summary!$C$5-1',
         f'=Summary!$C$5*(1+Summary!$C$7)^A{r}',
         f'=N{r}-S{r}', held, f'=Summary!$C$5*D{r}/Summary!$C$8', f'=Summary!$C$5*(1+Summary!$C$6*7*A{r}/365)']
    Y.append(row)
    for c in range(1,24): Y.cell(r,c).font=BLK
    for c in (1,2,3,4,11,12,21): Y.cell(r,c).font=BLUE
    for c in (13,):
        if n==1: Y.cell(r,c).font=GRN
    for c,fmt in ((3,'0.00'),(4,'0.00'),(5,PCT),(8,CUR),(9,CUR),(10,CUR),(11,CUR),(12,CUR),(13,CUR),(14,CUR),(15,CUR),(16,PCT2),(18,PCT),(19,CUR),(20,CUR),(22,CUR),(23,CUR)):
        Y.cell(r,c).number_format=fmt
    Y.cell(r,6).alignment=Alignment(wrap_text=True,vertical='top')
    if n in (26,52):
        for c in range(1,24): Y.cell(r,c).fill=SUB
nW=len(W)+1
Y.cell(nW+2,1,'Blue = recorded data from the run (prices, regime, share P&L, interest, shares held). Black = formulas. Green = link to Summary sheet.').font=Font(name=F,size=9,italic=True)
Y.cell(nW+3,1,'Share P&L = stock gain/loss that week from assignments, called-away sales and the Friday mark on shares held. Week 26 and Week 52 rows are shaded.').font=Font(name=F,size=9,italic=True)
Y.cell(nW+4,1,'Source: SimQ Run 2 books (run-002-state.json / run-002-ledger.md), simulated paper trading; regime is hidden during the week and revealed at Friday close.').font=Font(name=F,size=9,italic=True)
widths=[7,12,10,10,10,60,8,12,12,12,12,10,13,13,12,9,8,11,14,14,9,13,13]
for i,w in enumerate(widths,1): Y.column_dimensions[L(i)].width=w
Y.freeze_panes='B2'

# ---------------- Summary
S['A1']='SimQ Run 2: Weekly 1% Summary'; S['A1'].font=Font(name=F,size=14,bold=True)
S['A2']='Paper-trading simulation of cash-secured puts and covered calls on SimQ (a synthetic 3x leveraged index). Week-by-week detail is on the Weekly and Trades tabs.'
S['A2'].font=Font(name=F,size=9,italic=True)
S['A4']='Assumptions'; S['A4'].font=BOLD
ass=[('Starting cash ($)',7500,CUR0,'Run rules: identical start for every run'),
     ('Interest on idle cash (APY)',0.04,PCT,'Run rules: simple daily accrual, actual/365'),
     ('Weekly return target',0.01,PCT,'The "weekly 1%" goal being tested'),
     ('SimQ starting price ($)',63.80,'0.00','Run rules: SimQ at $63.80 on Week 1 open')]
for i,(lab,v,fmt,src) in enumerate(ass):
    r=5+i; S.cell(r,2,lab).font=BLK; c=S.cell(r,3,v); c.font=BLUE; c.number_format=fmt; S.cell(r,4,src).font=Font(name=F,size=9,italic=True)
S['C7'].fill=YEL
S['D7'].value='The "weekly 1%" goal being tested (change this to test another target)'

hr=11
S.cell(hr-1,1,'Results').font=BOLD
heads=['','Start','Week 26','Week 52','Start → Week 26','Week 26 → Week 52','Full year']
for j,h in enumerate(heads,1):
    c=S.cell(hr,j,h); c.font=HDR; c.fill=HFILL; c.alignment=Alignment(horizontal='center',wrap_text=True)
WE=f'Weekly!$N$2:$N${nW}'
def wk(col,n): return f'INDEX(Weekly!${col}$2:${col}${nW},{n})'
lines=[]
def add(label,start,w26,w52,a=None,b=None,full=None,fmt=CUR):
    lines.append((label,start,w26,w52,a,b,full,fmt))
add('Account value ($) (period cols = $ change)','=$C$5','='+wk('N',26),'='+wk('N',52),'=C{r}-B{r}','=D{r}-C{r}','=D{r}-B{r}')
add('Return vs prior point','',"=C{rp}/B{rp}-1","=D{rp}/B{rp}-1",'','','',PCT)
add('Avg weekly return (compounded)','','','','=(C{rv}/B{rv})^(1/26)-1','=(D{rv}/C{rv})^(1/26)-1','=(D{rv}/B{rv})^(1/52)-1',PCT2)
add('1% weekly target value ($)','=$C$5','=$C$5*(1+$C$7)^26','=$C$5*(1+$C$7)^52',None,None,None)
add('Ahead / behind 1% target ($)','','=C{rv}-C{r1}','=D{rv}-D{r1}',None,None,None)
add('SimQ price (period cols = % change)','=$C$8','='+wk('D',26),'='+wk('D',52),'=C{r}/B{r}-1','=D{r}/C{r}-1','=D{r}/B{r}-1','0.00')
add('Buy-and-hold SimQ value ($)','=$C$5','=$C$5*C{rs}/$C$8','=$C$5*D{rs}/$C$8',None,None,None)
add('Cash-only value at APY ($)','=$C$5','=$C$5*(1+$C$6*182/365)','=$C$5*(1+$C$6*364/365)',None,None,None)
add('Premium collected ($)','','','',f'=SUM(Weekly!H2:H27)',f'=SUM(Weekly!H28:H{nW})',f'=SUM(Weekly!H2:H{nW})')
add('Buy-backs paid ($)','','','',f'=SUM(Weekly!I2:I27)',f'=SUM(Weekly!I28:I{nW})',f'=SUM(Weekly!I2:I{nW})')
add('Options net ($)','','','',f'=SUM(Weekly!J2:J27)',f'=SUM(Weekly!J28:J{nW})',f'=SUM(Weekly!J2:J{nW})')
add('Share P&L ($)','','','',f'=SUM(Weekly!K2:K27)',f'=SUM(Weekly!K28:K{nW})',f'=SUM(Weekly!K2:K{nW})')
add('Interest ($)','','','',f'=SUM(Weekly!L2:L27)',f'=SUM(Weekly!L28:L{nW})',f'=SUM(Weekly!L2:L{nW})')
add('Total P&L ($)','','','','=SUM(E{p0}:E{p1})','=SUM(F{p0}:F{p1})','=SUM(G{p0}:G{p1})')
add('Trades opened','','','',f'=SUM(Weekly!G2:G27)',f'=SUM(Weekly!G28:G{nW})',f'=SUM(Weekly!G2:G{nW})','0')
add('Weeks that hit 1% or more','','','',f'=COUNTIF(Weekly!Q2:Q27,"Yes")',f'=COUNTIF(Weekly!Q28:Q{nW},"Yes")',f'=COUNTIF(Weekly!Q2:Q{nW},"Yes")','0')
add('Up weeks','','','',f'=COUNTIF(Weekly!O2:O27,">0")',f'=COUNTIF(Weekly!O28:O{nW},">0")',f'=COUNTIF(Weekly!O2:O{nW},">0")','0')
add('Down weeks','','','',f'=COUNTIF(Weekly!O2:O27,"<0")',f'=COUNTIF(Weekly!O28:O{nW},"<0")',f'=COUNTIF(Weekly!O2:O{nW},"<0")','0')
add('Best week %','','','',f'=MAX(Weekly!P2:P27)',f'=MAX(Weekly!P28:P{nW})',f'=MAX(Weekly!P2:P{nW})',PCT2)
add('Worst week %','','','',f'=MIN(Weekly!P2:P27)',f'=MIN(Weekly!P28:P{nW})',f'=MIN(Weekly!P2:P{nW})',PCT2)
add('Puts assigned','','','',f'=COUNTIFS(Trades!B2:B{nT},"<=26",Trades!L2:L{nT},"Assigned")',f'=COUNTIFS(Trades!B2:B{nT},">26",Trades!L2:L{nT},"Assigned")',f'=COUNTIF(Trades!L2:L{nT},"Assigned")','0')
add('Calls called away','','','',f'=COUNTIFS(Trades!B2:B{nT},"<=26",Trades!L2:L{nT},"Called away")',f'=COUNTIFS(Trades!B2:B{nT},">26",Trades!L2:L{nT},"Called away")',f'=COUNTIF(Trades!L2:L{nT},"Called away")','0')
rows={}
for i,l in enumerate(lines): rows[l[0]]=hr+1+i
rv=rows['Account value ($) (period cols = $ change)']; r1=rows['1% weekly target value ($)']; rs=rows['SimQ price (period cols = % change)']
p0=rows['Premium collected ($)']; p1=rows['Interest ($)']
for i,(label,a,b,c_,e,f,g,fmt) in enumerate(lines):
    r=hr+1+i; rp=rv
    ctx=dict(r=r,rp=rp,rv=rv,r1=r1,rs=rs,p0=p0+2,p1=p1)
    ctx['p0']=rows['Options net ($)']; # total = options net + shares + interest
    vals=[label,a,b,c_,e,f,g]
    for j,v in enumerate(vals,1):
        if v is None: v=''
        if isinstance(v,str) and '{' in v: v=v.format(**ctx)
        cell=S.cell(r,j,v); cell.font=BOLD if j==1 else BLK
        if j>1: cell.number_format=fmt
        cell.border=BOX
    if label.startswith('Account value') or label=='Total P&L ($)':
        for j in range(1,8): S.cell(r,j).font=BOLD
# SimQ period cols as percent
for j in (5,6,7): S.cell(rs,j).number_format=PCT
# Return row: start→W26 column E should be W26 vs start; fix row label semantics
r=rows['Return vs prior point']; S.cell(r,1,'Return since start'); S.cell(r,2,'')

# period changes for value row are $ changes; percent changes in a dedicated row
end=hr+len(lines)
S.cell(end+2,1,'How to read: Start/Week 26/Week 52 columns are snapshots; the last three columns are what happened during each period. Week 26 is the official cross-run comparison point; Week 52 is the full year.').font=Font(name=F,size=9,italic=True)
S.cell(end+3,1,'Benchmarks: buy-and-hold = $7,500 invested in SimQ at $63.80 and held. Cash-only = $7,500 left in cash at the APY. 1% target = starting cash compounded at the weekly target.').font=Font(name=F,size=9,italic=True)
S.cell(end+4,1,'Source: SimQ Run 2 books (run-002-state.json), simulated paper trading with simulated prices. Not real-market results.').font=Font(name=F,size=9,italic=True)
for i,w in enumerate([32,14,14,14,16,18,14],1): S.column_dimensions[L(i)].width=w
S.row_dimensions[hr].height=30
for row in S.iter_rows(min_row=4,max_row=9):
    for c in row:
        if c.font and c.font.name!=F: c.font=Font(name=F,size=10)

# ---------------- Run 1 vs Run 2 at Week 26
import json as _j
R1=_j.load(open('run1.json'))
cr=end+7
S.cell(cr-1,1,'Run 1 vs Run 2 at Week 26 (same start, same rules, different market)').font=BOLD
for j,h in enumerate(['Measure','Run 1 (Week 26)','Run 2 (Week 26)','Change'],1):
    c=S.cell(cr,j,h); c.font=HDR; c.fill=HFILL; c.alignment=Alignment(horizontal='center',wrap_text=True)
R=rows; rv=R['Account value ($) (period cols = $ change)']
cmp=[('Account value ($)',R1['end'],f'=C{rv}',CUR,'diff'),
     ('Return since start',None,f'=C{R["Return vs prior point"]}',PCT,'diff'),
     ('Avg weekly return (compounded)',None,f'=E{R["Avg weekly return (compounded)"]}',PCT2,'diff'),
     ('Ahead / behind 1% target ($)',None,f'=C{R["Ahead / behind 1% target ($)"]}',CUR,'diff'),
     ('SimQ move over 26 weeks',27.53,f'=E{R["SimQ price (period cols = % change)"]}',PCT,'none'),
     ('Premium collected ($)',R1['prem'],f'=E{R["Premium collected ($)"]}',CUR,'diff'),
     ('Buy-backs paid ($)',R1['btc'],f'=E{R["Buy-backs paid ($)"]}',CUR,'diff'),
     ('Options net ($)',None,f'=E{R["Options net ($)"]}',CUR,'diff'),
     ('Share P&L ($)',R1['shares'],f'=E{R["Share P&L ($)"]}',CUR,'diff'),
     ('Interest ($)',R1['interest'],f'=E{R["Interest ($)"]}',CUR,'diff'),
     ('Trades opened',R1['trades'],f'=E{R["Trades opened"]}','0','diff'),
     ('Weeks that hit 1% or more',R1['hit'],f'=E{R["Weeks that hit 1% or more"]}','0','diff'),
     ('Up weeks',R1['up'],f'=E{R["Up weeks"]}','0','diff'),
     ('Down weeks',R1['down'],f'=E{R["Down weeks"]}','0','diff'),
     ('Best week %',R1['best'],f'=E{R["Best week %"]}',PCT2,'diff'),
     ('Worst week %',R1['worst'],f'=E{R["Worst week %"]}',PCT2,'diff'),
     ('Puts assigned',R1['assigned'],f'=E{R["Puts assigned"]}','0','diff'),
     ('Calls called away',R1['called'],f'=E{R["Calls called away"]}','0','diff'),
     ('Shares still held at Week 26','100 (marked $27.53, cost $42.50)','0 (all cash)',None,'none')]
r1row={}
for i,(lab,v1,v2,fmt,kind) in enumerate(cmp):
    r=cr+1+i; r1row[lab]=r
    S.cell(r,1,lab).font=BOLD
    c1=S.cell(r,2,v1); c2=S.cell(r,3,v2); c2.font=GRN if str(v2).startswith('=') else BLK
    c1.font=BLUE
    if fmt: c1.number_format=fmt; c2.number_format=fmt
    if kind=='diff':
        c=S.cell(r,4,f'=C{r}-B{r}'); c.number_format=fmt; c.font=BLK
    for j in range(1,5): S.cell(r,j).border=BOX
a=r1row['Account value ($)']
S.cell(r1row['Return since start'],2,f'=B{a}/$C$5-1').font=BLK
S.cell(r1row['Avg weekly return (compounded)'],2,f'=(B{a}/$C$5)^(1/26)-1').font=BLK
S.cell(r1row['Ahead / behind 1% target ($)'],2,f'=B{a}-$C$5*(1+$C$7)^26').font=BLK
q=r1row['SimQ move over 26 weeks']; S.cell(q,2,f'=27.53/$C$8-1').font=BLK
S.cell(q,2).comment=Comment('Run 1 SimQ closed Week 26 at $27.53 (run-001-summary.txt).','Claude')
o=r1row['Options net ($)']; S.cell(o,2,f'=B{r1row["Premium collected ($)"]}-B{r1row["Buy-backs paid ($)"]}').font=BLK
ce=cr+len(cmp)
S.cell(ce+1,1,'Run 1 figures (blue) are from Run 1\'s books (run-001-state.json / run-001-summary.txt). Run 1 used older pricing (v1) through Week 12; Run 2 used v2 pricing with put skew from Week 1. Different seeds mean different markets: SimQ fell 57% in Run 1 and rose 44% in Run 2 over the same 26 weeks.').font=Font(name=F,size=9,italic=True)
S.cell(ce+2,1,'Run 1 share P&L = -$450 realized on lot 1 plus -$1,497 unrealized on lot 2, still held at Week 26.').font=Font(name=F,size=9,italic=True)

# ---------------- Chart
from openpyxl.chart import LineChart, Reference
ch=LineChart()
ch.title='Run 2 account vs the 1% weekly target'
ch.y_axis.title='Value ($)'; ch.x_axis.title='Week'
ch.height=9; ch.width=22
for col,name in ((14,None),(19,None),(22,None),(23,None)):
    ref=Reference(Y,min_col=col,min_row=1,max_row=nW)
    ch.add_data(ref,titles_from_data=True)
ch.set_categories(Reference(Y,min_col=1,min_row=2,max_row=nW))
colors=['1F3864','C00000','7F7F7F','70AD47']
for sr,colr in zip(ch.series,colors):
    sr.graphicalProperties.line.solidFill=colr; sr.graphicalProperties.line.width=22000; sr.smooth=False
ch.series[1].graphicalProperties.line.dashStyle='dash'
ch.y_axis.number_format='$#,##0'
ch.y_axis.delete=False; ch.x_axis.delete=False
ch.legend.position='b'
S.add_chart(ch,'I4')
S.cell(3,9,'Chart: account end value each week (blue), the 1% weekly target path (red dashed), buy-and-hold SimQ (grey), and cash-only (green). Data: Weekly tab columns N, S, V, W.').font=Font(name=F,size=9,italic=True)

wb.save('SimQ_Run2_Weekly_1pct.xlsx')
print(rows)
